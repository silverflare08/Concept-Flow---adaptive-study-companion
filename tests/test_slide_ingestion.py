from pathlib import Path

from pptx import Presentation

from studybuddy.ingestion.slides import ingest_slides


def test_slide_ingestion_preserves_slide_numbers(tmp_path: Path) -> None:
    deck_path = tmp_path / "graphs.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.slides[0].shapes.add_textbox(0, 0, 1000000, 500000).text_frame.text = "BFS uses a queue"
    presentation.slides[1].shapes.add_textbox(0, 0, 1000000, 500000).text_frame.text = "DFS uses a stack"
    presentation.save(deck_path)

    result = ingest_slides(
        deck_path,
        source_id="graphs-slides",
        source_name="Graph slides",
        source_uri="https://example.edu/graphs.pptx",
    )

    assert [segment.citation.slide for segment in result.segments] == [1, 2]
    assert result.segments[0].citation.target_uri().endswith("#page=1")
    assert "queue" in result.segments[0].text
