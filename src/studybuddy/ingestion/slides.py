"""PowerPoint ingestion with slide-cited text and visual evidence."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .models import CitationAnchor, EvidenceKind, SourceSegment, SourceType

ImageDescriber = Callable[[Path], str | None]


@dataclass(frozen=True)
class SlideIngestionResult:
    source_id: str
    source_sha256: str
    segments: list[SourceSegment]
    extracted_images: list[Path]


def _sha256(path: Path) -> str:
    hasher = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            hasher.update(block)
    return hasher.hexdigest()


def _text_from_slide(slide: object) -> str:
    paragraphs: list[str] = []
    for shape in slide.shapes:  # type: ignore[attr-defined]
        if getattr(shape, "has_text_frame", False):
            text = shape.text.strip()
            if text:
                paragraphs.append(text)
    return "\n".join(paragraphs)


def ingest_slides(
    path: str | Path,
    *,
    source_id: str,
    source_name: str,
    source_uri: str | None = None,
    image_output_dir: str | Path | None = None,
    describe_image: ImageDescriber | None = None,
) -> SlideIngestionResult:
    """Extract one text segment per slide and one optional segment per figure.

    A slide's text is intentionally not merged with descriptions of its figures.
    That distinction lets a retrieval result make clear whether a response is
    grounded in a diagram, OCR, or ordinary slide text.
    """
    deck_path = Path(path)
    if not deck_path.is_file():
        raise FileNotFoundError(f"Slide deck not found: {deck_path}")

    image_dir = Path(image_output_dir) if image_output_dir else deck_path.parent / "images"
    presentation = Presentation(str(deck_path))
    segments: list[SourceSegment] = []
    extracted_images: list[Path] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        anchor = CitationAnchor(
            source_id=source_id,
            source_name=source_name,
            source_type=SourceType.SLIDE,
            source_uri=source_uri,
            slide=slide_number,
        )
        text = _text_from_slide(slide)
        if text:
            segments.append(SourceSegment(text=text, citation=anchor, evidence_kind=EvidenceKind.TEXT))

        for image_index, shape in enumerate(slide.shapes, start=1):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            image_dir.mkdir(parents=True, exist_ok=True)
            extension = shape.image.ext.lower()
            output = image_dir / f"{source_id}-s{slide_number}-img{image_index}.{extension}"
            output.write_bytes(shape.image.blob)
            extracted_images.append(output)

            if describe_image is None:
                continue
            description = describe_image(output)
            if description and description.strip():
                segments.append(
                    SourceSegment(
                        text=description.strip(),
                        citation=anchor,
                        evidence_kind=EvidenceKind.FIGURE_DESCRIPTION,
                        image_path=str(output),
                    )
                )

    return SlideIngestionResult(
        source_id=source_id,
        source_sha256=_sha256(deck_path),
        segments=segments,
        extracted_images=extracted_images,
    )
